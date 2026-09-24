from __future__ import annotations

import json
import os
import sys
from pathlib import Path

from docx import Document
from pypdf import PdfReader
import xlrd


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKSPACE_ROOT = PROJECT_ROOT.parent
sys.path.insert(0, str(PROJECT_ROOT / "backend"))

from app.services.document_parser import parse_document  # noqa: E402


TEMPLATE_DIRS = (
    WORKSPACE_ROOT / "KSOCM_合同模板",
    WORKSPACE_ROOT / "POC合同模板（租赁）",
)


def independent_metadata(path: Path) -> dict:
    suffix = path.suffix.lower()
    result: dict = {}
    if suffix == ".pdf":
        reader = PdfReader(path)
        page_text = [(page.extract_text() or "").strip() for page in reader.pages]
        result.update(
            pages=len(reader.pages),
            independent_text_chars=sum(len(text) for text in page_text),
            pages_with_text=sum(bool(text) for text in page_text),
        )
    elif suffix == ".docx":
        document = Document(path)
        result.update(
            paragraphs=len(document.paragraphs),
            tables=len(document.tables),
            inline_shapes=len(document.inline_shapes),
            independent_text_chars=sum(len(paragraph.text) for paragraph in document.paragraphs)
            + sum(len(cell.text) for table in document.tables for row in table.rows for cell in row.cells),
        )
    elif suffix == ".xls":
        workbook = xlrd.open_workbook(path)
        result.update(
            sheets=workbook.nsheets,
            rows=sum(sheet.nrows for sheet in workbook.sheets()),
            independent_text_chars=sum(
                len(str(value))
                for sheet in workbook.sheets()
                for row_index in range(sheet.nrows)
                for value in sheet.row_values(row_index)
                if value not in (None, "")
            ),
        )
    return result


def main() -> None:
    rows = []
    requested = {value.casefold() for value in sys.argv[1:]}
    for directory in TEMPLATE_DIRS:
        for path in sorted(directory.iterdir()):
            if not path.is_file():
                continue
            if requested and path.name.casefold() not in requested:
                continue
            parsed = parse_document(path, path.name)
            row = {
                "folder": directory.name,
                "file": path.name,
                "extension": path.suffix.lower(),
                "size_bytes": path.stat().st_size,
                "parse_status": parsed.status,
                "parse_message": parsed.message,
                "parsed_chars": len(parsed.text),
                "parsed_lines": len(parsed.text.splitlines()),
                "fields": parsed.fields,
                "sample": parsed.text[:240].replace("\n", " "),
            }
            try:
                row.update(independent_metadata(path))
            except Exception as exc:  # noqa: BLE001
                row["independent_error"] = f"{type(exc).__name__}: {exc}"
            rows.append(row)
    payload = json.dumps(rows, ensure_ascii=False, indent=2)
    output_path = os.environ.get("TEMPLATE_INVENTORY_OUTPUT", "").strip()
    if output_path:
        destination = Path(output_path)
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_text(payload, encoding="utf-8")
    print(payload)


if __name__ == "__main__":
    main()
