from __future__ import annotations

import base64
import hashlib
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from functools import lru_cache
from io import BytesIO
from pathlib import Path
from urllib.parse import urlparse
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from .annotation_service import extract_document_annotations


WORD_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@dataclass
class ParseResult:
    text: str
    status: str = "已解析"
    message: str = ""
    fields: dict = field(default_factory=dict)
    annotations: list[dict] = field(default_factory=list)
    source_map: list[dict] = field(default_factory=list)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _run_text_command(command: list[str], timeout: int = 120) -> str:
    completed = subprocess.run(command, capture_output=True, check=False, timeout=timeout)
    for encoding in ("utf-8", "gb18030", "utf-16"):
        try:
            return completed.stdout.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return completed.stdout.decode("utf-8", errors="ignore").strip()


@lru_cache(maxsize=1)
def _rapid_ocr():
    from rapidocr_onnxruntime import RapidOCR

    return RapidOCR(det_limit_side_len=2000)


def _deduplicate_lines(lines: list[str]) -> list[str]:
    result: list[str] = []
    for value in lines:
        text = re.sub(r"\s+", " ", value).strip()
        if not text or text in result[-5:]:
            continue
        result.append(text)
    return result


def _ocr_pil_image(image: Image.Image) -> str:
    image = image.convert("RGB")
    if image.width < 80 or image.height < 80:
        return ""
    if image.width > 1800:
        ratio = 1800 / image.width
        image = image.resize((1800, max(1, round(image.height * ratio))), Image.Resampling.LANCZOS)
    # Small Chinese glyphs disappear when a whole A4 page is downscaled by the
    # OCR detector. Horizontal bands preserve the original character size.
    band_height = 900
    overlap = 30
    starts = [0] if image.height <= band_height else list(range(0, image.height, band_height - overlap))
    lines: list[str] = []
    engine = _rapid_ocr()
    for top in starts:
        bottom = min(image.height, top + band_height)
        if top and bottom - top < 200:
            break
        detected, _ = engine(image.crop((0, top, image.width, bottom)))
        if detected:
            detected = sorted(
                detected,
                key=lambda item: (
                    min(point[1] for point in item[0]),
                    min(point[0] for point in item[0]),
                ),
            )
            lines.extend(str(item[1]) for item in detected if len(item) >= 2)
        if bottom >= image.height:
            break
    return "\n".join(_deduplicate_lines(lines))


def _ocr_image_bytes(data: bytes) -> str:
    try:
        with Image.open(BytesIO(data)) as image:
            return _ocr_pil_image(image)
    except Exception:
        return ""


def _word_blocks(xml_data: bytes, tag: str, label: str) -> list[str]:
    from xml.etree import ElementTree

    root = ElementTree.fromstring(xml_data)
    values: list[str] = []
    for block in root.iter(f"{WORD_NS}{tag}"):
        text = "".join(
            node.text or ""
            for node in block.iter()
            if node.tag in {f"{WORD_NS}t", f"{WORD_NS}delText"}
        ).strip()
        if text:
            values.append(f"[{label}] {text}")
    return values


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _effective_xml_text(node) -> str:
    """Extract effective Word text while excluding tracked deletions.

    Reading ``Document.paragraphs`` first and ``Document.tables`` afterwards
    destroys the source order. Contract retrieval needs paragraphs and tables
    in their original sequence, so the body XML is the authority here.
    """
    name = _local_name(node.tag)
    if name in {"del", "moveFrom"}:
        return ""
    if name == "t":
        return node.text or ""
    if name == "tab":
        return "\t"
    if name in {"br", "cr"}:
        return "\n"
    return "".join(_effective_xml_text(child) for child in list(node))


def _rendered_break_count(node) -> int:
    rendered = sum(1 for item in node.iter() if _local_name(item.tag) == "lastRenderedPageBreak")
    if rendered:
        return rendered
    return sum(
        1
        for item in node.iter()
        if _local_name(item.tag) == "br" and any(_local_name(key) == "type" and value == "page" for key, value in item.attrib.items())
    )


def _append_source_line(lines: list[tuple[str, int, str]], value: str, page_no: int, block_type: str) -> None:
    text = re.sub(r"[ \t]+", " ", value or "").strip()
    text = re.sub(r"\n{3,}", "\n\n", text)
    if text:
        lines.append((text, max(1, page_no), block_type))


def _docx_ordered_lines(path: Path) -> tuple[list[tuple[str, int, str]], list[tuple[str, bytes]]]:
    from xml.etree import ElementTree

    lines: list[tuple[str, int, str]] = []
    media: list[tuple[str, bytes]] = []
    with ZipFile(path) as package:
        names = set(package.namelist())
        root = ElementTree.fromstring(package.read("word/document.xml"))
        body = next((item for item in root.iter() if _local_name(item.tag) == "body"), None)
        if body is None:
            raise ValueError("Word 正文结构不存在")
        page_no = 1
        for block in list(body):
            block_name = _local_name(block.tag)
            if block_name == "p":
                _append_source_line(lines, _effective_xml_text(block), page_no, "paragraph")
                page_no += _rendered_break_count(block)
            elif block_name == "tbl":
                for row in (item for item in list(block) if _local_name(item.tag) == "tr"):
                    row_page = page_no
                    values: list[str] = []
                    for cell in (item for item in list(row) if _local_name(item.tag) == "tc"):
                        values.append(re.sub(r"\s+", " ", _effective_xml_text(cell)).strip())
                    if any(values):
                        _append_source_line(lines, " | ".join(values), row_page, "table_row")
                    page_no += _rendered_break_count(row)
        for name in sorted(item for item in names if item.startswith("word/media/")):
            media.append((Path(name).name, package.read(name)))
    return lines, media


def _rendered_docx_pages(path: Path) -> list[str]:
    """Render Word to PDF when LibreOffice is available for authoritative pages."""
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    try:
        with tempfile.TemporaryDirectory(prefix="contract-pages-") as output_dir:
            pdf_path = Path(output_dir) / f"{path.stem}.pdf"
            if executable:
                completed = subprocess.run(
                    [executable, "--headless", "--convert-to", "pdf", "--outdir", output_dir, str(path)],
                    capture_output=True,
                    check=False,
                    timeout=180,
                )
                if completed.returncode != 0:
                    return []
            else:
                powershell = shutil.which("powershell")
                if not powershell:
                    return []
                source_value = str(path.resolve()).replace("'", "''")
                target_value = str(pdf_path.resolve()).replace("'", "''")
                script = f"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
try {{
  $document = $word.Documents.Open('{source_value}', $false, $true)
  $document.ExportAsFixedFormat('{target_value}', 17)
  $document.Close($false)
}} finally {{
  $word.Quit()
}}
"""
                encoded = base64.b64encode(script.encode("utf-16le")).decode("ascii")
                completed = subprocess.run(
                    [powershell, "-NoProfile", "-NonInteractive", "-EncodedCommand", encoded],
                    capture_output=True,
                    check=False,
                    timeout=180,
                )
                if completed.returncode != 0:
                    return []
            if not pdf_path.exists():
                return []
            reader = PdfReader(pdf_path)
            return [(page.extract_text() or "").strip() for page in reader.pages]
    except Exception:
        return []


def _compact(value: str) -> str:
    # Table extraction inserts synthetic pipes while PDF text layers use
    # columns/spaces. Remove layout separators before cross-render matching.
    return re.sub(r"[\s|¦]+", "", value or "")


def _apply_rendered_pages(source_map: list[dict], pages: list[str]) -> None:
    if not pages:
        return
    compact_pages = [_compact(page) for page in pages]
    last_page = 1
    for item in source_map:
        needle = _compact(item.get("text", ""))
        if len(needle) < 8:
            continue
        # Long table rows sometimes wrap differently in PDF extraction. Match
        # a stable prefix and suffix before falling back to Word page-breaks.
        probes = [needle]
        if len(needle) > 80:
            probes.extend((needle[:60], needle[-60:]))
        matches = [
            page_index + 1
            for page_index, page in enumerate(compact_pages)
            if any(len(probe) >= 8 and probe in page for probe in probes)
        ]
        if matches:
            selected = next((value for value in matches if value >= last_page), matches[0])
            item["page_no"] = selected
            item["page_method"] = "rendered_pdf"
            last_page = selected


def _parse_docx_structured(path: Path) -> tuple[str, list[dict], str]:
    ordered, media = _docx_ordered_lines(path)
    text_parts: list[str] = []
    source_map: list[dict] = []
    cursor = 0
    for value, page_no, block_type in ordered:
        if text_parts:
            text_parts.append("\n")
            cursor += 1
        start = cursor
        text_parts.append(value)
        cursor += len(value)
        source_map.append({
            "start_offset": start,
            "end_offset": cursor,
            "page_no": page_no,
            "page_method": "word_rendered_break",
            "block_type": block_type,
            "text": value,
        })
    for name, data in media:
        image_text = _ocr_image_bytes(data)
        if len(image_text) < 30:
            continue
        value = f"[文档图片OCR：{name}]\n{image_text}"
        if text_parts:
            text_parts.append("\n")
            cursor += 1
        start = cursor
        text_parts.append(value)
        cursor += len(value)
        source_map.append({
            "start_offset": start,
            "end_offset": cursor,
            "page_no": None,
            "page_method": "embedded_image_ocr",
            "block_type": "image_ocr",
            "text": value,
        })
    _apply_rendered_pages(source_map, _rendered_docx_pages(path))
    page_methods = {item.get("page_method") for item in source_map if item.get("page_no")}
    message = "已按Word原始段落/表格顺序解析"
    if "rendered_pdf" in page_methods:
        message += "并通过渲染页校准页码"
    else:
        message += "；页码来自Word最近一次分页标记，建议部署LibreOffice以获得渲染页校准"
    return "".join(text_parts), source_map, message


def _parse_docx(path: Path) -> str:
    return _parse_docx_structured(path)[0]


def _pdf_source_map(path: Path, text: str) -> list[dict]:
    """Bind extracted PDF text to pages when the text layer is available."""
    try:
        reader = PdfReader(path)
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
    except Exception:
        pages = []
    result: list[dict] = []
    cursor = 0
    for page_no, page_text in enumerate(pages, start=1):
        if not page_text:
            continue
        start = text.find(page_text, cursor)
        if start < 0:
            continue
        end = start + len(page_text)
        result.append({
            "start_offset": start,
            "end_offset": end,
            "page_no": page_no,
            "page_method": "pdf_text_layer",
            "block_type": "page",
            "text": page_text,
        })
        cursor = end
    if result:
        return result
    marker = re.compile(r"\[第(\d+)页\]\s*")
    matches = list(marker.finditer(text))
    for index, match in enumerate(matches):
        start = match.end()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        result.append({
            "start_offset": start,
            "end_offset": end,
            "page_no": int(match.group(1)),
            "page_method": "ocr_page_marker",
            "block_type": "page",
            "text": text[start:end].strip(),
        })
    return result


def _is_readable_contract_text(text: str) -> bool:
    if len(text.strip()) < 120:
        return False
    nonspace = [character for character in text if not character.isspace()]
    if not nonspace:
        return False
    controls = sum(ord(character) < 32 and character not in "\n\r\t" for character in text)
    cjk = sum("\u4e00" <= character <= "\u9fff" for character in nonspace)
    ascii_alnum = sum(character.isascii() and character.isalnum() for character in nonspace)
    if controls > max(5, len(text) * 0.005):
        return False
    return cjk / len(nonspace) >= 0.08 or ascii_alnum / len(nonspace) >= 0.45


def _ocr_pdf_pages(path: Path, page_indexes: list[int]) -> dict[int, str]:
    import fitz

    output: dict[int, str] = {}
    with fitz.open(path) as document:
        for page_index in page_indexes:
            page = document.load_page(page_index)
            pixmap = page.get_pixmap(dpi=150, alpha=False, annots=True)
            with Image.open(BytesIO(pixmap.tobytes("png"))) as image:
                output[page_index] = _ocr_pil_image(image)
    return output


def _pdf_annotation_text(reader: PdfReader) -> list[str]:
    lines: list[str] = []
    for page_number, page in enumerate(reader.pages, start=1):
        for reference in page.get("/Annots", []):
            try:
                annotation = reference.get_object()
                contents = str(annotation.get("/Contents", "")).strip()
                if contents:
                    subject = str(annotation.get("/Subject", "")).strip()
                    label = subject or str(annotation.get("/Subtype", "PDF批注")).strip("/")
                    lines.append(f"[PDF批注 第{page_number}页 {label}] {contents}")
            except Exception:
                continue
    return lines


def _parse_pdf(path: Path) -> tuple[str, str]:
    messages: list[str] = []
    try:
        reader = PdfReader(path)
        page_texts = [(page.extract_text() or "").strip() for page in reader.pages]
    except Exception as exc:
        reader = None
        page_texts = []
        messages.append(f"PDF文本层读取失败：{type(exc).__name__}")
    text = "\n\n".join(page_texts).strip()
    if _is_readable_contract_text(text):
        return text, "；".join(messages)

    executable = shutil.which("pdftotext")
    if executable:
        extracted = _run_text_command([executable, "-layout", str(path), "-"])
        if _is_readable_contract_text(extracted):
            return extracted, "；".join(messages)

    try:
        page_count = len(reader.pages) if reader else 0
        if not page_count:
            import fitz

            with fitz.open(path) as document:
                page_count = document.page_count
        ocr_pages = _ocr_pdf_pages(path, list(range(page_count)))
        text = "\n\n".join(
            f"[第{index + 1}页]\n{ocr_pages.get(index, '')}" for index in range(page_count)
        ).strip()
        messages.append(f"已对{page_count}页执行本地中文OCR")
    except Exception as exc:
        messages.append(f"PDF OCR失败：{type(exc).__name__}: {exc}")
    if not _is_readable_contract_text(text):
        messages.append("PDF未识别出足够的可读中文文本")
    return text, "；".join(messages)


def _parse_xlsx(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in workbook.worksheets:
        lines.append(f"【工作表：{sheet.title}】")
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value not in (None, "")]
            if values:
                lines.append(" | ".join(values))
    return "\n".join(lines)


def _parse_xls(path: Path) -> str:
    import xlrd

    workbook = xlrd.open_workbook(path)
    lines: list[str] = []
    for sheet in workbook.sheets():
        lines.append(f"【工作表：{sheet.name}】")
        for row_index in range(sheet.nrows):
            values = [str(value).strip() for value in sheet.row_values(row_index) if str(value).strip()]
            if values:
                lines.append(" | ".join(values))
    return "\n".join(lines)


def _parse_image(path: Path) -> tuple[str, str]:
    text = ""
    executable = shutil.which("tesseract")
    if executable:
        text = _run_text_command([executable, str(path), "stdout", "-l", "chi_sim+eng"], timeout=300)
    if not _is_readable_contract_text(text):
        with path.open("rb") as stream:
            text = _ocr_image_bytes(stream.read())
    return text, "本地中文OCR识别完成" if text else "OCR未识别出文字"


def _parse_legacy_doc(path: Path) -> tuple[str, str]:
    executable = shutil.which("antiword")
    if executable:
        text = _run_text_command([executable, str(path)])
        if text.strip():
            return text, "旧版Word已通过antiword解析"

    # Windows local deployments normally have Word even when antiword is not
    # installed. Convert invisibly to DOCX and reuse the same rich parser.
    powershell = shutil.which("powershell")
    if powershell:
        with tempfile.TemporaryDirectory(prefix="contract-doc-") as directory:
            converted = Path(directory) / f"{path.stem}.docx"
            source_value = str(path.resolve()).replace("'", "''")
            target_value = str(converted.resolve()).replace("'", "''")
            script = f"""
$ErrorActionPreference = 'Stop'
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
try {{
  $document = $word.Documents.Open('{source_value}', $false, $true)
  $document.SaveAs2('{target_value}', 16)
  $document.Close($false)
}} finally {{
  $word.Quit()
}}
"""
            encoded = base64.b64encode(script.encode("utf-16le")).decode("ascii")
            completed = subprocess.run(
                [powershell, "-NoProfile", "-NonInteractive", "-EncodedCommand", encoded],
                capture_output=True,
                check=False,
                timeout=180,
            )
            if completed.returncode == 0 and converted.exists():
                return _parse_docx(converted), "旧版Word已通过本机Word转换解析"
    return "", "旧版Word解析失败：需要antiword，Windows本机可使用Microsoft Word转换"


def extract_fields(text: str, file_name: str) -> dict:
    compact = re.sub(r"[ \t]+", " ", text)
    fields: dict[str, str | float] = {}
    contract_numbers = re.findall(
        r"(?:合同编号|合同号|协议号|订单号码|订单号|Document\s*Number|Contract\s*No\.?)\s*[：:]?\s*([A-Za-z0-9_\-/（）()\u4e00-\u9fff]{4,50})",
        compact,
        re.IGNORECASE,
    )
    for candidate in contract_numbers:
        if re.search(r"\d", candidate) and candidate.casefold() not in {"contractno", "documentnumber"}:
            fields["contract_no"] = candidate
            break

    title = ""
    clean_stem = re.sub(r"^\d+[.\-]\s*", "", Path(file_name).stem).strip()
    if "一般性条款" in clean_stem:
        title = re.sub(r"\s*脱敏\s*$", "", clean_stem).strip()
    elif re.search(r"(?:^|[.\-_ ])PO(?:[.\-_ ]|$)", file_name, re.IGNORECASE) or re.search(
        r"purchase\s*order", compact[:1200], re.IGNORECASE
    ):
        title = "采购订单"
    known_title = re.search(
        r"(林德\s*[（(]中国[）)]\s*叉车有限公司设备(?:销售|租赁)[^\n]{0,24}(?:一般性条款(?:和条件)?))",
        compact[:2200],
    )
    if known_title:
        title = re.sub(r"\s+", "", known_title.group(1))
    for raw_line in text.splitlines()[:20]:
        if title:
            break
        line = re.sub(r"\s+", " ", raw_line).strip(" _【】[]")
        if not 3 <= len(line) <= 70 or not re.search(r"(?:合同|协议|采购订单|一般性条款)$", line):
            continue
        if re.search(r"[，。；]", line) or re.search(
            r"^(?:若|如|本合同|除|双方|工作表|[0-9]+(?:\.[0-9]+)?)", line
        ):
            continue
        title = line
        break
    if not title:
        match = re.search(
            r"([\u4e00-\u9fffA-Za-z0-9（）()·\- ]{2,60}(?:采购合同|销售合同|服务合同|租赁[（(]框架[）)]合同|租赁框架合同|租赁合同|租赁协议|合同书|一般性条款))",
            compact[:2200],
        )
        title = match.group(1).strip() if match else Path(file_name).stem
    fields["name"] = title[:120]

    party_a = re.search(
        r"(?:甲方|买方|承租方|采购方)\s*(?:[（(][^）)\n]{0,12}[）)])?\s*[：:]\s*[【\[]?([^\n|】\]]{2,100})",
        compact[:6000],
    )
    if party_a:
        customer = party_a.group(1).strip(" _【】[]，,。;；")
        if not re.match(r"^(?:[一二三四五六七八九十]+、|[0-9]+[）).、]|若|如)", customer):
            fields["customer"] = customer
    if "customer" not in fields:
        po_buyer = re.search(r"^(VOSS\s+Automotive\s+Components[^\n]{0,80})", text, re.IGNORECASE | re.MULTILINE)
        if po_buyer:
            fields["customer"] = po_buyer.group(1).strip()

    amount_patterns = [
        r"(?:合同总价|合同价款|价税合计|总金额|合同金额)[^\n0-9]{0,40}(?:人民币|RMB|CNY|[￥])?\s*([0-9][0-9,]*(?:\.\d{1,2})?)\s*(万元|元|人民币|RMB|CNY)",
    ]
    for pattern in amount_patterns:
        match = re.search(pattern, compact, re.IGNORECASE)
        if match:
            value = float(match.group(1).replace(",", ""))
            unit = match.group(2) if match.lastindex and match.lastindex >= 2 else ""
            value = value * 10000 if unit == "万元" else value
            if value > 0:
                fields["amount"] = value
                break

    lowered = compact.lower()
    classification_text = f"{file_name}\n{title}\n{compact[:2600]}"
    classification_head = classification_text[:1200]
    file_and_title = f"{file_name}\n{title}"
    if "租赁" in file_and_title or "租金" in file_and_title:
        fields["contract_type"] = "租赁"
    elif re.search(r"(?:设备销售|采购订单|采购合同|销售合同)", file_and_title):
        fields["contract_type"] = "采购"
    elif re.search(r"(?:设备销售|采购订单|采购合同|销售合同)", classification_head):
        fields["contract_type"] = "采购"
    elif "租赁" in classification_head or "租金" in classification_head:
        fields["contract_type"] = "租赁"
    else:
        # The rebuilt POC has two business lanes only. Purchase orders,
        # equipment sales and supporting service procurements all enter the
        # procurement lane; equipment-use contracts enter the lease lane.
        fields["contract_type"] = "采购"

    url = re.search(r"(?:https?://|www\.)[^\s，。；;）)\]】]+", lowered)
    if url:
        candidate = url.group(0)
        if candidate.startswith("www."):
            candidate = "https://" + candidate
        if urlparse(candidate).netloc:
            fields["website_terms_url"] = candidate
    if "website_terms_url" not in fields:
        # OCR commonly drops the first dot ("wwwVOSS.net") or turns it into
        # whitespace ("www voss.de"). Rebuild only an explicit www-domain so
        # ordinary dotted text cannot become a false URL.
        ocr_url = re.search(
            r"\bwww(?:\.|\s+)?([a-z0-9][a-z0-9-]{1,62})\s*\.\s*([a-z]{2,})(/[^\s，。；;）)\]\u3011]*)?",
            lowered,
        )
        if ocr_url:
            path = ocr_url.group(3) or ""
            fields["website_terms_url"] = f"https://www.{ocr_url.group(1)}.{ocr_url.group(2)}{path}"
    date_matches = re.findall(r"(20\d{2})[年./-](\d{1,2})[月./-](\d{1,2})", compact)
    if date_matches:
        fields["dates"] = [
            "-".join([year, month.zfill(2), day.zfill(2)]) for year, month, day in date_matches[:8]
        ]
    rent = re.search(
        r"(?:月租金|每月租金)[^\n0-9]{0,30}([0-9][0-9,]*(?:\.\d{1,2})?)\s*(?:元|万元)",
        compact,
    )
    if rent:
        fields["monthly_rent"] = rent.group(1)
    payment = re.search(r"(?:付款方式(?:与条件)?|付款条件|支付方式)\s*[：:]?[^。；\n]{0,180}", compact)
    if payment:
        fields["payment_terms"] = payment.group(0)
    return fields


def parse_document(path: Path, original_name: str) -> ParseResult:
    suffix = path.suffix.lower()
    source_map: list[dict] = []
    try:
        if suffix == ".docx":
            text, source_map, message = _parse_docx_structured(path)
        elif suffix == ".doc":
            text, message = _parse_legacy_doc(path)
        elif suffix == ".pdf":
            text, message = _parse_pdf(path)
            source_map = _pdf_source_map(path, text)
        elif suffix in {".xlsx", ".xlsm"}:
            text, message = _parse_xlsx(path), ""
        elif suffix == ".xls":
            text, message = _parse_xls(path), ""
        elif suffix in {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}:
            text, message = _parse_image(path)
        elif suffix in {".txt", ".md", ".csv"}:
            text, message = path.read_text(encoding="utf-8", errors="ignore"), ""
        else:
            text, message = "", "文件已安全保存；此格式需通过外部解析器或多模态模型处理"
    except Exception as exc:
        return ParseResult(text="", status="解析待处理", message=f"{type(exc).__name__}: {exc}", fields={}, annotations=[], source_map=[])
    status = "已解析" if _is_readable_contract_text(text) else "解析待处理"
    annotations = [item.to_dict() for item in extract_document_annotations(path)]
    if annotations:
        message = "；".join(part for part in (message, f"识别到{len(annotations)}条批注或修订，已单独归档来源") if part)
    for item in source_map:
        item["source_file"] = original_name
        # Keeping full block text in both extracted_text and the JSON map is
        # unnecessary and can double database size. Offsets are authoritative.
        item.pop("text", None)
    return ParseResult(
        text=text.strip(),
        status=status,
        message=message,
        fields=extract_fields(text, original_name),
        annotations=annotations,
        source_map=source_map,
    )
