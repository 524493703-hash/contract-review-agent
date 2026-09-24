from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from zipfile import ZipFile

import fitz
from docx import Document
from docx.enum.text import WD_COLOR_INDEX


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{W_NS}}}"

PDF_AI_HIGHLIGHT_COLOR = (0.31, 0.80, 0.84)
DOCX_AI_HIGHLIGHT_COLOR = WD_COLOR_INDEX.TURQUOISE


@dataclass
class DetectedAnnotation:
    external_id: str = ""
    annotation_type: str = "批注"
    author_name: str = ""
    page_no: int | None = None
    anchor_text: str = ""
    comment_text: str = ""
    comment_date: str = ""
    location: dict = field(default_factory=dict)

    def to_dict(self) -> dict:
        return asdict(self)


def _xml_text(element) -> str:
    return "".join(
        node.text or ""
        for node in element.iter()
        if node.tag in {f"{W}t", f"{W}delText"}
    ).strip()


def _docx_annotations(path: Path) -> list[DetectedAnnotation]:
    from xml.etree import ElementTree

    rows: list[DetectedAnnotation] = []
    with ZipFile(path) as package:
        names = set(package.namelist())
        comment_meta: dict[str, dict] = {}
        if "word/comments.xml" in names:
            root = ElementTree.fromstring(package.read("word/comments.xml"))
            for item in root.iter(f"{W}comment"):
                comment_id = item.attrib.get(f"{W}id", "")
                comment_meta[comment_id] = {
                    "author": item.attrib.get(f"{W}author", ""),
                    "date": item.attrib.get(f"{W}date", ""),
                    "text": _xml_text(item),
                }

        anchor_parts: dict[str, list[str]] = {comment_id: [] for comment_id in comment_meta}
        story_names = [
            name for name in names
            if name == "word/document.xml" or re.match(r"word/(?:header|footer)\d+\.xml$", name)
        ]
        for story_name in story_names:
            root = ElementTree.fromstring(package.read(story_name))
            for paragraph in root.iter(f"{W}p"):
                active: list[str] = []

                def visit(node) -> None:
                    tag = node.tag
                    comment_id = node.attrib.get(f"{W}id", "")
                    if tag == f"{W}commentRangeStart" and comment_id:
                        active.append(comment_id)
                        return
                    if tag == f"{W}commentRangeEnd" and comment_id:
                        if comment_id in active:
                            active.remove(comment_id)
                        return
                    if tag in {f"{W}t", f"{W}delText"} and node.text:
                        for active_id in active:
                            anchor_parts.setdefault(active_id, []).append(node.text)
                    for child in list(node):
                        visit(child)

                visit(paragraph)

            for tag, label in (("ins", "修订插入"), ("del", "修订删除")):
                for change in root.iter(f"{W}{tag}"):
                    text = _xml_text(change)
                    if not text:
                        continue
                    rows.append(DetectedAnnotation(
                        external_id=change.attrib.get(f"{W}id", ""),
                        annotation_type=label,
                        author_name=change.attrib.get(f"{W}author", ""),
                        anchor_text=text,
                        comment_text=text,
                        comment_date=change.attrib.get(f"{W}date", ""),
                        location={"story": story_name},
                    ))

        for comment_id, meta in comment_meta.items():
            rows.append(DetectedAnnotation(
                external_id=comment_id,
                annotation_type="Word批注",
                author_name=meta["author"],
                anchor_text=re.sub(r"\s+", " ", "".join(anchor_parts.get(comment_id, []))).strip(),
                comment_text=meta["text"],
                comment_date=meta["date"],
                location={"story": "word/document.xml"},
            ))
    return rows


def _pdf_annotations(path: Path) -> list[DetectedAnnotation]:
    rows: list[DetectedAnnotation] = []

    def point_pair(point) -> list[float]:
        """PyMuPDF versions expose annotation vertices as Points or tuples."""
        x = point.x if hasattr(point, "x") else point[0]
        y = point.y if hasattr(point, "y") else point[1]
        return [round(float(x), 3), round(float(y), 3)]

    with fitz.open(path) as document:
        for page_index, page in enumerate(document):
            annotation = page.first_annot
            while annotation:
                annotation_type = annotation.type[1]
                if annotation_type != "Popup":
                    info = annotation.info or {}
                    rect = annotation.rect
                    anchor = re.sub(r"\s+", " ", page.get_textbox(rect)).strip()
                    rows.append(DetectedAnnotation(
                        external_id=str(info.get("id") or annotation.xref),
                        annotation_type=f"PDF{annotation_type}",
                        author_name=str(info.get("title") or ""),
                        page_no=page_index + 1,
                        anchor_text=anchor,
                        comment_text=str(info.get("content") or ""),
                        comment_date=str(info.get("modDate") or info.get("creationDate") or ""),
                        location={
                            "rect": [round(value, 3) for value in (rect.x0, rect.y0, rect.x1, rect.y1)],
                            "vertices": [point_pair(point) for point in (annotation.vertices or [])],
                        },
                    ))
                annotation = annotation.next
    return rows


def extract_document_annotations(path: Path) -> list[DetectedAnnotation]:
    suffix = path.suffix.lower()
    try:
        if suffix == ".docx":
            return _docx_annotations(path)
        if suffix == ".pdf":
            return _pdf_annotations(path)
    except Exception:
        return []
    return []


def resolve_annotation_source(
    annotation: DetectedAnnotation,
    document_source: str,
    source_department: str,
    users: list,
    identities: list,
) -> dict:
    author = re.sub(r"\s+", " ", annotation.author_name or "").strip()
    folded = author.casefold()
    for identity in identities:
        if folded and folded == str(identity.alias).strip().casefold():
            return {
                "source_kind": identity.source_kind,
                "source_department": identity.department,
                "source_confidence": 99,
                "source_basis": f"批注作者命中身份映射：{identity.alias}",
                "author_name": identity.display_name or author,
            }
    for user in users:
        aliases = {str(user.name).casefold(), str(user.username).casefold()}
        if user.email:
            aliases.add(str(user.email).casefold())
            aliases.add(str(user.email).split("@", 1)[0].casefold())
        if folded and folded in aliases:
            return {
                "source_kind": "同事",
                "source_department": user.department,
                "source_confidence": 98,
                "source_basis": f"批注作者匹配员工账号：{user.name}",
                "author_name": user.name,
            }
    department_tokens = {
        "法务": "法务部", "legal": "法务部", "财务": "财务部", "finance": "财务部",
        "销售": "销售部", "sales": "销售部", "合同": "合同管理部", "abu": "ABU",
        "服务": "服务部", "service": "服务部",
    }
    for token, department in department_tokens.items():
        if token.casefold() in folded:
            return {
                "source_kind": "部门",
                "source_department": department,
                "source_confidence": 88,
                "source_basis": f"根据批注作者名称识别部门关键词：{token}",
                "author_name": author,
            }
    if document_source == "客户":
        return {
            "source_kind": "客户",
            "source_department": "客户",
            "source_confidence": 84 if author else 72,
            "source_basis": "根据本轮文件来源及批注作者元数据综合判断",
            "author_name": author or "客户（未署名）",
        }
    if source_department or document_source.endswith("部") or document_source in {"ABU", "其他内部"}:
        department = source_department or document_source
        return {
            "source_kind": "部门",
            "source_department": department,
            "source_confidence": 80 if author else 68,
            "source_basis": "根据本轮内部文件来源判断，作者未命中员工映射",
            "author_name": author or f"{department}（未署名）",
        }
    return {
        "source_kind": "未识别",
        "source_department": "",
        "source_confidence": 30 if author else 10,
        "source_basis": "未找到作者身份映射，请人工确认",
        "author_name": author or "未署名",
    }


def _evidence_candidates(evidence: str) -> list[str]:
    value = re.sub(r"【(?:命中\d+|另有\d+处同类命中)】", " ", str(evidence or ""))
    value = value.replace("…", " ")
    pieces = [re.sub(r"\s+", " ", item).strip() for item in re.split(r"[\n。；！？]", value)]
    candidates: list[str] = []
    for piece in sorted((item for item in pieces if len(re.sub(r"\s+", "", item)) >= 8), key=len, reverse=True):
        if len(piece) > 90:
            for start in range(0, min(len(piece), 180), 45):
                part = piece[start:start + 60].strip()
                if len(re.sub(r"\s+", "", part)) >= 12:
                    candidates.append(part)
        candidates.append(piece[:120])
    return list(dict.fromkeys(candidates))[:12]


def _comment_body(suggestion, round_no: int = 1) -> str:
    return (
        f"【R{round_no} · {suggestion.risk_level}风险】{suggestion.title}\n"
        f"建议：{suggestion.suggested_text}\n"
        f"依据：{suggestion.basis}\n"
        f"谈判重点：{suggestion.negotiation_focus}"
    )[:3500]


def _pdf_span_rects(page, compact_needle: str) -> list[fitz.Rect]:
    if len(compact_needle) < 8:
        return []
    chars: list[str] = []
    boxes: list[tuple[float, float, float, float]] = []
    for block in page.get_text("dict").get("blocks", []):
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                bbox = tuple(span.get("bbox", (0, 0, 0, 0)))
                for char in str(span.get("text", "")):
                    if not char.isspace():
                        chars.append(char)
                        boxes.append(bbox)
    compact_page = "".join(chars)
    start = compact_page.find(compact_needle)
    if start < 0:
        return []
    unique: list[fitz.Rect] = []
    for bbox in boxes[start:start + len(compact_needle)]:
        rect = fitz.Rect(bbox)
        if not unique or tuple(round(v, 2) for v in unique[-1]) != tuple(round(v, 2) for v in rect):
            unique.append(rect)
    return unique[:12]


def _annotate_pdf(source: Path, suggestions: list, output: Path, round_no: int = 1) -> dict:
    anchored = 0
    fallback = 0
    with fitz.open(source) as document:
        for suggestion in suggestions:
            placed = False
            for page in document:
                rects: list[fitz.Rect] = []
                for candidate in _evidence_candidates(suggestion.anchor_text or suggestion.original_text):
                    rects = page.search_for(candidate, quads=False)
                    if not rects:
                        rects = _pdf_span_rects(page, re.sub(r"\s+", "", candidate))
                    if rects:
                        break
                if not rects:
                    continue
                for index, rect in enumerate(rects[:8]):
                    annotation = page.add_highlight_annot(rect)
                    annotation.set_colors(stroke=PDF_AI_HIGHLIGHT_COLOR)
                    if index == 0:
                        annotation.set_info(
                            title=f"契析 AI审核 · R{round_no}",
                            subject=f"R{round_no} · {suggestion.risk_level}风险 · {suggestion.category}",
                            content=_comment_body(suggestion, round_no),
                        )
                    annotation.update()
                suggestion.page_no = page.number + 1
                anchored += 1
                placed = True
                break
            if not placed:
                page = document[0]
                point = fitz.Point(max(18, page.rect.width - 34), 34 + fallback * 18)
                note = page.add_text_annot(point, _comment_body(suggestion, round_no), icon="Comment")
                note.set_info(title=f"契析 AI审核 · R{round_no}", subject=f"R{round_no} · 未能精确定位 · {suggestion.category}")
                note.update()
                fallback += 1
        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(output, garbage=4, deflate=True)
    return {"format": "pdf", "anchored": anchored, "fallback": fallback, "total": len(suggestions)}


def _iter_docx_paragraphs(document: Document):
    seen: set[int] = set()

    def emit(paragraphs):
        for paragraph in paragraphs:
            marker = id(paragraph._p)
            if marker not in seen:
                seen.add(marker)
                yield paragraph

    yield from emit(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from emit(cell.paragraphs)
    for section in document.sections:
        for container in (section.header, section.footer):
            yield from emit(container.paragraphs)


def _matching_runs(paragraph, candidate: str):
    raw = "".join(run.text for run in paragraph.runs)
    compact_chars: list[str] = []
    raw_offsets: list[int] = []
    for index, char in enumerate(raw):
        if not char.isspace():
            compact_chars.append(char)
            raw_offsets.append(index)
    needle = re.sub(r"\s+", "", candidate)
    start = "".join(compact_chars).find(needle)
    if start < 0:
        return []
    raw_start = raw_offsets[start]
    raw_end = raw_offsets[start + len(needle) - 1] + 1
    matches = []
    cursor = 0
    for run in paragraph.runs:
        next_cursor = cursor + len(run.text)
        if cursor < raw_end and next_cursor > raw_start and run.text:
            matches.append(run)
        cursor = next_cursor
    return matches


def _annotate_docx(source: Path, suggestions: list, output: Path, round_no: int = 1) -> dict:
    document = Document(source)
    paragraphs = list(_iter_docx_paragraphs(document))
    anchored = 0
    fallback = 0
    for suggestion in suggestions:
        target_runs = []
        for candidate in _evidence_candidates(suggestion.anchor_text or suggestion.original_text):
            for paragraph in paragraphs:
                target_runs = _matching_runs(paragraph, candidate)
                if target_runs:
                    break
            if target_runs:
                break
        if target_runs:
            for run in target_runs:
                run.font.highlight_color = DOCX_AI_HIGHLIGHT_COLOR
            document.add_comment(target_runs, _comment_body(suggestion, round_no), author=f"契析 AI审核 · R{round_no}", initials=f"R{round_no}")
            anchored += 1
        else:
            fallback += 1
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return {"format": "docx", "anchored": anchored, "fallback": fallback, "total": len(suggestions)}


def annotate_source_document(source: Path, suggestions: list, output: Path, round_no: int = 1) -> dict:
    suffix = source.suffix.lower()
    if suffix == ".pdf":
        return _annotate_pdf(source, suggestions, output, round_no)
    if suffix == ".docx":
        return _annotate_docx(source, suggestions, output, round_no)
    return {"format": suffix.lstrip(".") or "unknown", "anchored": 0, "fallback": len(suggestions), "total": len(suggestions), "unsupported": True}
