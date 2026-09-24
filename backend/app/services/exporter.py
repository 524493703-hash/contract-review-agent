from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt, RGBColor

from .review_engine import apply_accepted_suggestions


ACCEPTED_DECISIONS = {"接纳", "修改后接纳"}


def _add_contract_text(doc: Document, text: str) -> None:
    doc.add_heading("合同正文", level=1)
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if line:
            doc.add_paragraph(line)


def build_effective_contract_docx(contract_name: str, effective_text: str, output_path: Path) -> Path:
    """Persist accepted working text as the source document for the next round."""
    doc = Document()
    doc.styles["Normal"].font.name = "Microsoft YaHei"
    doc.styles["Normal"].font.size = Pt(10.5)
    title = doc.add_heading(contract_name, level=0)
    title.runs[0].font.name = "Microsoft YaHei"
    subtitle = doc.add_paragraph("系统修订稿 · 已应用上一轮接纳的审核建议")
    subtitle.runs[0].font.color.rgb = RGBColor(80, 93, 112)
    _add_contract_text(doc, effective_text)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


def _apply_accepted_changes(source_text: str, suggestions: list) -> tuple[str, list]:
    rendered, _, unmatched_ids = apply_accepted_suggestions(source_text, suggestions)
    unmatched_id_set = set(unmatched_ids)
    unmatched = [suggestion for suggestion in suggestions if suggestion.id in unmatched_id_set]
    return rendered, unmatched


def build_review_docx(
    contract_name: str,
    source_text: str,
    suggestions: list,
    mode: str,
    output_path: Path,
    internal_notes: list[str] | None = None,
    focused_findings: list[dict] | None = None,
) -> Path:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"].font.size = Pt(10.5)
    title = doc.add_heading(contract_name, level=0)
    title.runs[0].font.name = "Microsoft YaHei"
    subtitle = doc.add_paragraph(f"审核输出版本：{ {'clean': '客户清洁版', 'redline': '带修改痕迹版', 'internal': '内部审核版'}.get(mode, mode) }")
    subtitle.runs[0].font.color.rgb = RGBColor(80, 93, 112)

    accepted = [item for item in suggestions if item.decision in ACCEPTED_DECISIONS]
    if mode == "clean":
        clean_text, unmatched = _apply_accepted_changes(source_text, suggestions)
        _add_contract_text(doc, clean_text)
        if unmatched:
            doc.add_heading("双方已确认的补充修改", level=1)
            for suggestion in unmatched:
                doc.add_paragraph(suggestion.suggested_text)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)
        return output_path

    _add_contract_text(doc, source_text)
    visible_suggestions = suggestions if mode == "internal" else accepted
    doc.add_heading("已接纳修改痕迹" if mode == "redline" else "内部审核意见", level=1)
    if not visible_suggestions:
        doc.add_paragraph("暂无已接纳的修改。" if mode == "redline" else "暂无审核意见。")
    for index, suggestion in enumerate(visible_suggestions, start=1):
        doc.add_heading(f"{index}. {suggestion.title}", level=1)
        if mode == "redline":
            original = doc.add_paragraph()
            run = original.add_run(suggestion.original_text)
            run.font.strike = True
            run.font.color.rgb = RGBColor(164, 46, 46)
            proposed = doc.add_paragraph()
            run = proposed.add_run(suggestion.suggested_text)
            run.font.color.rgb = RGBColor(14, 116, 144)
            run.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
        else:
            doc.add_paragraph(suggestion.suggested_text)
        if mode == "internal":
            table = doc.add_table(rows=3, cols=2)
            table.style = "Light Shading Accent 1"
            values = (("风险等级", suggestion.risk_level), ("修改依据", suggestion.basis), ("谈判重点", suggestion.negotiation_focus))
            for row, (label, value) in zip(table.rows, values):
                row.cells[0].text = label
                row.cells[1].text = value
    if mode == "internal" and internal_notes:
        doc.add_heading("内部协同与谈判备注", level=1)
        for note in internal_notes:
            if note.strip():
                doc.add_paragraph(note.strip())
    if mode == "internal" and focused_findings:
        doc.add_heading("Focused Checks 与具体条款依据", level=1)
        for item in focused_findings:
            if item.get("status") == "MET":
                continue
            doc.add_heading(
                f"{item.get('check_code', '')} · {item.get('check_name', '')} · {item.get('status', '')}",
                level=2,
            )
            doc.add_paragraph(str(item.get("reason") or ""))
            evidence = item.get("evidence") or []
            if evidence:
                doc.add_paragraph(f"问题合同原文：{evidence[0].get('quote', '')}")
            citations = item.get("baseline_citations") or []
            if citations:
                doc.add_paragraph("公司基准依据：")
            for citation in citations:
                page_label = f"，第{citation.get('page_no')}页" if citation.get("page_no") else ""
                doc.add_paragraph(
                    f"《{citation.get('document_title', '')}》{citation.get('document_version', '')}版，"
                    f"文件 {citation.get('source_file', '')}，第{citation.get('clause_no', '')}条{page_label}："
                    f"{citation.get('clause_text', '')}"
                )
            legal_citations = item.get("legal_citations") or []
            if legal_citations:
                doc.add_paragraph("法律依据：")
            for citation in legal_citations:
                doc.add_paragraph(
                    f"{citation.get('title', '')} {citation.get('article_no', '')}（{citation.get('document_no', '')}）："
                    f"{citation.get('summary', '')}；官方来源：{citation.get('source_url', '')}"
                )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path
