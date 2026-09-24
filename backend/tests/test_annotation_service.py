from __future__ import annotations

from types import SimpleNamespace

import fitz
import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from app.services.annotation_service import annotate_source_document, extract_document_annotations, resolve_annotation_source
from app.services.preview_service import AI_HIGHLIGHT_HEX, locate_pdf_risks, prepare_pdf_preview, render_docx_preview


def test_pdf_comment_author_detection_and_original_file_annotation(tmp_path):
    source = tmp_path / "customer-commented.pdf"
    document = fitz.open()
    page = document.new_page()
    sentence = "Payment must be completed within 60 days after acceptance."
    page.insert_text((72, 90), sentence, fontsize=12)
    customer = page.add_highlight_annot(page.search_for("within 60 days"))
    customer.set_info(title="Customer A", content="Please keep the 60-day payment term.")
    customer.update()
    document.save(source)
    document.close()

    detected = extract_document_annotations(source)
    assert len(detected) == 1
    assert detected[0].author_name == "Customer A"
    resolved = resolve_annotation_source(detected[0], "客户", "客户", [], [])
    assert resolved["source_kind"] == "客户"
    assert resolved["source_confidence"] >= 80

    suggestion = SimpleNamespace(
        id="risk-payment-1",
        risk_level="中",
        title="付款账期偏长",
        category="付款条件",
        original_text=sentence,
        anchor_text=sentence,
        suggested_text="Payment must be completed within 30 days after acceptance.",
        basis="Procurement and lease payment policy.",
        negotiation_focus="Keep the unsecured payment term within 45 days.",
        page_no=None,
    )
    output = tmp_path / "customer-commented-ai.pdf"
    result = annotate_source_document(source, [suggestion], output, round_no=2)
    assert result == {"format": "pdf", "anchored": 1, "fallback": 0, "total": 1}
    with fitz.open(output) as annotated:
        infos = []
        page = annotated[0]
        annotation = page.first_annot
        while annotation:
            if annotation.type[1] != "Popup":
                infos.append(annotation.info)
            annotation = annotation.next
    assert any(info.get("title") == "Customer A" for info in infos)
    assert any(info.get("title") == "契析 AI审核 · R2" and "【R2 ·" in info.get("content", "") and "谈判重点" in info.get("content", "") for info in infos)
    assert prepare_pdf_preview(output).exists()
    locations = locate_pdf_risks(output, [suggestion])
    assert locations["risk-payment-1"]
    assert locations["risk-payment-1"][0]["page_no"] == 1
    assert 0 <= locations["risk-payment-1"][0]["x"] <= 100


def test_docx_web_preview_keeps_original_comments_and_uses_uniform_ai_cyan(tmp_path):
    source = tmp_path / "round-four.docx"
    document = Document()
    original = document.add_paragraph().add_run("客户原文件黄色高亮，并带有客户批注。")
    original.font.highlight_color = WD_COLOR_INDEX.YELLOW
    document.add_comment([original], "客户要求保留该表述。", author="客户王经理", initials="客户")
    sentence = "租期届满后视同自动续租。"
    document.add_paragraph(sentence)
    document.save(source)

    suggestion = SimpleNamespace(
        id="risk-renewal-1",
        risk_level="高",
        title="自动续租风险",
        category="租赁期限",
        original_text=sentence,
        anchor_text=sentence,
        suggested_text="租期届满即终止，续租应另行书面确认。",
        basis="租赁合同一般性条款。",
        negotiation_focus="避免默示续租。",
        page_no=None,
    )
    output = tmp_path / "round-four-annotated.docx"
    result = annotate_source_document(source, [suggestion], output, round_no=4)
    assert result["anchored"] == 1

    rendered = render_docx_preview(output, [suggestion])
    assert "doc-highlight-yellow" in rendered["html"]
    assert "doc-highlight-cyan" in rendered["html"]
    assert "doc-ai-highlight" in rendered["html"]
    assert 'data-risk-id="risk-renewal-1"' in rendered["html"]
    assert "客户王经理" in rendered["html"]
    assert rendered["embedded_comment_count"] == 2
    assert AI_HIGHLIGHT_HEX == "#4fcbd6"


def test_docx_preview_links_cross_cell_risk_without_word_comment(tmp_path):
    source = tmp_path / "cross-cell-risk.docx"
    document = Document()
    table = document.add_table(rows=1, cols=4)
    for cell, value in zip(
        table.rows[0].cells,
        ["电动叉车 2T", "租赁期不超过一年", "未税价 2654.87", "含税价 30000"],
    ):
        cell.text = value
    document.save(source)
    suggestion = SimpleNamespace(
        id="risk-tax-math-1",
        title="未税价与含税价算术不一致",
        anchor_text="电动叉车 2T | 租赁期不超过一年 | 未税价 2654.87 | 含税价 30000",
        original_text="",
        suggested_text="",
    )

    rendered = render_docx_preview(source, [suggestion])

    assert 'data-risk-id="risk-tax-math-1"' in rendered["html"]
    assert 'class="doc-ai-risk-fallback"' in rendered["html"]


def test_legacy_pdf_multiline_ai_highlight_is_fully_normalized_without_recoloring_customer_mark(tmp_path):
    source = tmp_path / "legacy-ai.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 90), "AI first line", fontsize=12)
    page.insert_text((72, 120), "AI continuation", fontsize=12)
    page.insert_text((72, 150), "Customer highlight", fontsize=12)
    first = page.add_highlight_annot(page.search_for("AI first line"))
    first.set_colors(stroke=(1.0, 0.76, 0.28))
    first.set_info(title="契析 AI审核", content="legacy")
    first.update()
    continuation = page.add_highlight_annot(page.search_for("AI continuation"))
    continuation.set_colors(stroke=(1.0, 0.76, 0.28))
    continuation.update()
    customer = page.add_highlight_annot(page.search_for("Customer highlight"))
    customer.set_colors(stroke=(1.0, 0.9, 0.1))
    customer.set_info(title="Customer")
    customer.update()
    document.save(source)
    document.close()

    preview = prepare_pdf_preview(source)
    with fitz.open(preview) as normalized:
        annotations = []
        page = normalized[0]
        annotation = page.first_annot
        while annotation:
            if annotation.type[1] == "Highlight":
                annotations.append((annotation.info.get("title"), annotation.colors["stroke"]))
            annotation = annotation.next
    assert annotations[0][1] == pytest.approx([0.31, 0.80, 0.84], abs=0.01)
    assert annotations[1][1] == pytest.approx([0.31, 0.80, 0.84], abs=0.01)
    assert annotations[2][0] == "Customer"
    assert annotations[2][1] == pytest.approx([1.0, 0.9, 0.1], abs=0.01)
