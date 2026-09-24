from __future__ import annotations

import os
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from fastapi.testclient import TestClient
from docx import Document
from sqlalchemy import event


TEST_DB = Path(os.environ["DATABASE_URL"].removeprefix("sqlite:///"))
os.environ["LLM_API_KEY"] = ""
os.environ["JWT_SECRET"] = "test-only-secret"

from app.main import app  # noqa: E402
from app.database import engine  # noqa: E402


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as test_client:
        yield test_client
    engine.dispose()
    TEST_DB.unlink(missing_ok=True)


def login(client: TestClient, username: str) -> dict[str, str]:
    response = client.post("/api/auth/login", json={"username": username, "password": "Poc@2026"})
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def test_health_login_and_permissions(client: TestClient):
    assert client.get("/api/health").json()["status"] == "ok"
    assert client.get("/api/contracts").status_code == 401
    sales = login(client, "sales.owner")
    contracts = client.get("/api/contracts", headers=sales)
    assert contracts.status_code == 200
    assert len(contracts.json()) >= 3
    contract_id = contracts.json()[0]["id"]
    assert client.post(f"/api/contracts/{contract_id}/review", headers=sales).status_code == 403


def test_login_bootstrap_and_ready_probe(client: TestClient):
    response = client.post("/api/auth/login", json={"username": "legal.reviewer", "password": "Poc@2026"})
    assert response.status_code == 200
    bootstrap = response.json()["bootstrap"]
    assert bootstrap["profile"]["username"] == "legal.reviewer"
    assert bootstrap["contracts"]
    assert bootstrap["users"]
    assert client.get("/api/ready").json()["status"] == "ok"
    headers = {"Authorization": f"Bearer {response.json()['access_token']}"}
    playbook = client.get("/api/playbooks/active", headers=headers)
    assert playbook.status_code == 200
    assert playbook.json()["coverage"]["status"] == "COMPLETE"
    assert len(playbook.json()["checks"]) == 49
    atomic_hit = client.get("/api/knowledge/structured", params={"q": "1000小时"}, headers=headers)
    assert atomic_hit.status_code == 200
    clause_11_4 = next(item for item in atomic_hit.json()["clauses"] if item["clause_no"] == "11.4")
    assert clause_11_4["atoms"]
    assert clause_11_4["focused_checks"]

    statements: list[str] = []
    def capture_statement(_conn, _cursor, statement, _parameters, _context, _executemany):
        statements.append(statement)
    event.listen(engine, "before_cursor_execute", capture_statement)
    try:
        assert client.get("/api/contracts", headers=headers).status_code == 200
    finally:
        event.remove(engine, "before_cursor_execute", capture_statement)
    selects = [statement for statement in statements if statement.lstrip().upper().startswith("SELECT")]
    assert len(selects) == 1, "合同列表必须保持单条聚合查询，禁止重新引入 N+1"


def test_upload_review_decide_and_export(client: TestClient, monkeypatch: pytest.MonkeyPatch):
    # Keep one end-to-end compatibility path for deployments that temporarily
    # opt back into visible legacy findings during migration.
    monkeypatch.setattr("app.main.settings.legacy_rule_mode", "active")
    legal = login(client, "legal.reviewer")
    def docx_payload(text: str, with_customer_comment: bool = False) -> bytes:
        document = Document()
        paragraph = document.add_paragraph(text)
        if with_customer_comment:
            document.add_comment(paragraph.runs, "客户坚持60日账期，请保留。", author="客户王经理", initials="客户")
        stream = BytesIO()
        document.save(stream)
        return stream.getvalue()

    payload = (
        "设备租赁合同\n租期届满后继续使用且出租方无异议的，视同自动续租。"
        "验收后60日支付。甲方有权单方解除合同。"
        "租赁设备的型号、数量和使用地点以设备清单为准，双方应按合同约定办理交付、验收、付款、维修保养、设备返还和争议解决事项。"
        "承租方应妥善保管设备并遵守操作手册和安全管理要求，发生事故时立即通知出租方。"
    )
    upload = client.post(
        "/api/contracts/upload",
        headers=legal,
        data={"customer": "自动化测试客户", "project": "POC端到端测试", "contract_type": "租赁", "document_source": "客户"},
        files=[("files", ("测试租赁合同-R1.docx", docx_payload(payload), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )
    assert upload.status_code == 200, upload.text
    contract_id = upload.json()[0]["contract"]["id"]
    review = client.post(f"/api/contracts/{contract_id}/review", headers=legal)
    assert review.status_code == 200, review.text
    assert review.json()["stats"]["total"] >= 2
    assert review.json()["legacy_shadow"]["mode"] == "active"
    assert review.json()["legacy_shadow"]["candidate_count"] == 0
    assert review.json()["focused_review"]["stats"]["TOTAL"] == 49
    focused = client.get(f"/api/review-runs/{review.json()['review_run_id']}/focused-checks", headers=legal)
    assert focused.status_code == 200, focused.text
    renewal = next(item for item in focused.json()["results"] if item["check_code"] == "L-43")
    assert renewal["status"] in {"MET", "PARTIAL"}
    if renewal["evidence"]:
        assert renewal["evidence"][0]["source_file"].endswith(".docx")
        assert "page_method" in renewal["evidence"][0]
    assert any(item["clause_no"] == "13.1" for item in renewal["baseline_citations"])
    assert all(item["source_file"].endswith("2025年01版.pdf") for item in renewal["baseline_citations"])
    assert all(item["clause_text"] for item in renewal["baseline_citations"])
    safety = next(item for item in focused.json()["results"] if item["check_code"] == "L-39")
    assert safety["legal_rag_required"] is True
    assert safety["legal_citations"]
    assert all(item["source_url"].startswith("https://") for item in safety["legal_citations"])
    legal_library = client.get("/api/legal-authorities?as_of=2026-09-10", headers=legal)
    assert legal_library.status_code == 200
    assert len(legal_library.json()["items"]) == 11
    assert review.json()["annotated_available"] is True
    first_annotated = client.get(
        f"/api/contracts/{contract_id}/review-runs/{review.json()['review_run_id']}/annotated",
        headers=legal,
    )
    assert first_annotated.status_code == 200
    with ZipFile(BytesIO(first_annotated.content)) as package:
        assert "word/comments.xml" in package.namelist()
        assert b"w:highlight" in package.read("word/document.xml")
        comments_xml = package.read("word/comments.xml").decode("utf-8")
        assert "依据：" in comments_xml
        assert "谈判重点" in comments_xml

    detail = client.get(f"/api/contracts/{contract_id}", headers=legal).json()
    first_version = detail["versions"][0]
    first_preview = client.get(
        f"/api/contracts/{contract_id}/versions/{first_version['id']}/preview",
        headers=legal,
    )
    assert first_preview.status_code == 200, first_preview.text
    preview_data = first_preview.json()
    assert preview_data["format"] == "docx"
    assert preview_data["is_annotated"] is True
    assert preview_data["round_color"] == "#4fcbd6"
    assert "doc-ai-highlight" in preview_data["html"]
    assert "data-risk-id=" in preview_data["html"]
    assert any(item["kind"] == "AI审核" and item["round_no"] == 1 for item in preview_data["annotations"])
    assert all("locations" in item for item in preview_data["annotations"])
    preview_file = client.get(
        f"/api/contracts/{contract_id}/versions/{first_version['id']}/preview/file",
        headers=legal,
    )
    assert preview_file.status_code == 200
    assert preview_file.content[:2] == b"PK"
    suggestion = next(item for item in detail["review_runs"][0]["suggestions"] if item["rule_code"] == "TERMINATION_ONE_SIDED")
    assert suggestion["action_type"] == "替换条款"
    assert suggestion["occurrence_no"] == 1
    assert suggestion["risk_key"]
    assert len(suggestion["revision_history"]) == 1
    assert not suggestion["suggested_text"].startswith(("建议", "删除", "修改", "调整"))

    transfer_target = next(item for item in detail["review_runs"][0]["suggestions"] if item["rule_code"] == "PAYMENT_LONG")
    finance_user = next(item for item in client.get("/api/auth/demo-users").json()["users"] if item["username"] == "finance.expert")
    transfer = client.post(
        f"/api/contracts/{contract_id}/consultations",
        headers=legal,
        json={"suggestion_id": transfer_target["id"], "assignee_id": finance_user["id"], "question": "请复核60日账期是否可以接受。", "due_hours": 24},
    )
    assert transfer.status_code == 200, transfer.text
    assert transfer.json()["status"] == "待审核"
    assert transfer.json()["target_department"] == "财务部"
    assert transfer.json()["requester_id"]
    duplicate_transfer = client.post(
        f"/api/contracts/{contract_id}/consultations",
        headers=legal,
        json={"suggestion_id": transfer_target["id"], "assignee_id": finance_user["id"], "question": "重复转批测试。"},
    )
    assert duplicate_transfer.status_code == 409
    blocked_decision = client.patch(
        f"/api/suggestions/{transfer_target['id']}",
        headers=legal,
        json={"decision": "接纳", "suggested_text": transfer_target["suggested_text"], "note": "不应越过转批"},
    )
    assert blocked_decision.status_code == 409
    assert client.patch(
        f"/api/consultations/{transfer.json()['id']}/answer",
        headers=legal,
        json={"review_decision": "建议修改", "answer": "请缩短账期。"},
    ).status_code == 403
    finance = login(client, "finance.expert")
    transfer_answer = client.patch(
        f"/api/consultations/{transfer.json()['id']}/answer",
        headers=finance,
        json={"review_decision": "建议修改", "answer": "建议账期不超过45日，并保留逾期付款责任。"},
    )
    assert transfer_answer.status_code == 200, transfer_answer.text
    assert transfer_answer.json()["status"] == "已审核"
    transferred_detail = client.get(f"/api/contracts/{contract_id}", headers=legal).json()
    stored_transfer = next(item for item in transferred_detail["consultations"] if item["id"] == transfer.json()["id"])
    assert stored_transfer["suggestion_id"] == transfer_target["id"]
    assert stored_transfer["review_decision"] == "建议修改"
    assert stored_transfer["target_department"] == "财务部"

    cancel_target = suggestion
    cancel_transfer = client.post(
        f"/api/contracts/{contract_id}/consultations",
        headers=legal,
        json={"suggestion_id": cancel_target["id"], "assignee_id": finance_user["id"], "question": "测试误转批后撤回。"},
    )
    assert cancel_transfer.status_code == 200
    cancelled = client.patch(f"/api/consultations/{cancel_transfer.json()['id']}/cancel", headers=legal)
    assert cancelled.status_code == 200
    assert cancelled.json()["status"] == "已撤回"

    placeholder = client.patch(
        f"/api/suggestions/{suggestion['id']}",
        headers=legal,
        json={"decision": "修改后接纳", "suggested_text": "应该XXX", "note": "占位测试"},
    )
    assert placeholder.status_code == 400, placeholder.text
    unrelated_clause = client.patch(
        f"/api/suggestions/{suggestion['id']}",
        headers=legal,
        json={
            "decision": "修改后接纳",
            "suggested_text": "因本合同引起的争议，任何一方均可向被告所在地有管辖权的人民法院提起诉讼。",
            "note": "关联性测试",
        },
    )
    assert unrelated_clause.status_code == 400, unrelated_clause.text
    assert "关联" in unrelated_clause.json()["detail"]

    def fake_redraft(**kwargs):
        instruction = kwargs["instruction"]
        clause = (
            "甲方仅在乙方发生重大违约且收到书面催告后四十五日内仍未纠正时方可解除本合同。"
            if "四十五" in instruction
            else "甲方仅在乙方发生重大违约且收到书面催告后三十日内仍未纠正时方可解除本合同。"
        )
        return clause, "测试大模型"

    monkeypatch.setattr("app.main.redraft_replacement_clause", fake_redraft)
    second_draft = client.post(
        f"/api/suggestions/{suggestion['id']}/redraft",
        headers=legal,
        json={"instruction": "将单方解除改为重大违约并给予三十日补救期"},
    )
    assert second_draft.status_code == 200, second_draft.text
    assert second_draft.json()["draft_version"] == 2
    assert len(second_draft.json()["revision_history"]) == 2
    third_draft = client.post(
        f"/api/suggestions/{suggestion['id']}/redraft",
        headers=legal,
        json={"instruction": "把补救期限改成四十五日"},
    )
    assert third_draft.status_code == 200, third_draft.text
    assert third_draft.json()["draft_version"] == 3
    assert len(third_draft.json()["revision_history"]) == 3
    accepted_text = third_draft.json()["suggested_text"]
    decision = client.patch(
        f"/api/suggestions/{suggestion['id']}",
        headers=legal,
        json={"decision": "修改后接纳", "suggested_text": accepted_text, "note": "接纳第三稿"},
    )
    assert decision.status_code == 200, decision.text
    assert decision.json()["decision"] == "修改后接纳"
    assert len(decision.json()["revision_history"]) == 3

    missing_round_upload = client.post(f"/api/contracts/{contract_id}/review", headers=legal)
    assert missing_round_upload.status_code == 409
    round_two_text = f"设备租赁合同\n{accepted_text}验收后60日支付。甲方有权单方解除合同。"
    round_two_upload = client.post(
        "/api/contracts/upload",
        headers=legal,
        data={"contract_id": contract_id, "review_round": "2", "contract_type": "租赁", "document_source": "客户"},
        files=[("files", ("测试租赁合同-R2.docx", docx_payload(round_two_text, with_customer_comment=True), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )
    assert round_two_upload.status_code == 200, round_two_upload.text
    assert round_two_upload.json()[0]["comment_count"] == 1
    second_review = client.post(f"/api/contracts/{contract_id}/review", headers=legal)
    assert second_review.status_code == 200, second_review.text
    assert second_review.json()["round_no"] == 2
    assert second_review.json()["inherited_accepted_opinions"] == 1
    second_detail = client.get(f"/api/contracts/{contract_id}", headers=legal).json()
    assert len(second_detail["versions"]) == 2
    assert second_detail["versions"][-1]["review_round_no"] == 2
    assert second_detail["versions"][-1]["document_source"] == "客户"
    assert second_detail["document_comments"][0]["source_kind"] == "客户"
    assert second_detail["document_comments"][0]["author_name"] == "客户王经理"
    second_preview = client.get(
        f"/api/contracts/{contract_id}/versions/{second_detail['versions'][-1]['id']}/preview",
        headers=legal,
    ).json()
    focused_risk_ids = {
        item["id"]
        for item in second_detail["focused_review"]["results"]
        if item["status"] in {"UNMET", "PARTIAL", "NOT_MENTIONED"}
        and item["technical_status"] == "verified"
    }
    expected_annotation_kinds = {"原文件批注", "AI审核"}
    if focused_risk_ids:
        expected_annotation_kinds.add("专项检查")
    assert {item["kind"] for item in second_preview["annotations"]} == expected_annotation_kinds
    assert all(item["round_no"] == 2 for item in second_preview["annotations"])
    assert all(item["highlight_color"] == "#4fcbd6" for item in second_preview["annotations"] if item["kind"] == "AI审核")
    focused_preview_ids = {
        item["id"] for item in second_preview["annotations"] if item["kind"] == "专项检查"
    }
    assert focused_preview_ids == focused_risk_ids
    for item in second_detail["focused_review"]["results"]:
        if item["status"] == "MET":
            continue
        assert item["negotiation_focus"]
        assert item["redline"]["proposed_text"]
    second_codes = {item["rule_code"] for item in second_detail["review_runs"][0]["suggestions"]}
    assert "LEASE_AUTO_RENEW" not in second_codes
    assert "PAYMENT_LONG" in second_codes
    payment = next(item for item in second_detail["review_runs"][0]["suggestions"] if item["rule_code"] == "PAYMENT_LONG")
    rejected = client.patch(
        f"/api/suggestions/{payment['id']}",
        headers=legal,
        json={"decision": "拒绝", "suggested_text": payment["suggested_text"], "note": "业务批准保留60日账期"},
    )
    assert rejected.status_code == 200, rejected.text
    round_three_upload = client.post(
        "/api/contracts/upload",
        headers=legal,
        data={"contract_id": contract_id, "review_round": "3", "contract_type": "租赁", "document_source": "法务部", "source_department": "法务部"},
        files=[("files", ("测试租赁合同-R3.docx", docx_payload(round_two_text), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
    )
    assert round_three_upload.status_code == 200, round_three_upload.text
    third_review = client.post(f"/api/contracts/{contract_id}/review", headers=legal)
    assert third_review.status_code == 200, third_review.text
    assert third_review.json()["round_no"] == 3
    third_detail = client.get(f"/api/contracts/{contract_id}", headers=legal).json()
    third_codes = {item["rule_code"] for item in third_detail["review_runs"][0]["suggestions"]}
    assert "PAYMENT_LONG" not in third_codes
    assert "TERMINATION_ONE_SIDED" in third_codes
    assert len(third_detail["versions"]) == 3
    assert third_detail["cumulative_review"]["rounds_completed"] == 3
    assert third_detail["cumulative_review"]["total_opinions"] >= len(third_detail["review_runs"][0]["suggestions"])
    generated = client.get(
        f"/api/contracts/{contract_id}/versions/{third_detail['versions'][-1]['id']}/download",
        headers=legal,
    )
    generated_document = Document(BytesIO(generated.content))
    generated_text = "\n".join(paragraph.text for paragraph in generated_document.paragraphs)
    assert accepted_text in generated_text

    internal_secret = "内部底线：若客户不恢复标准条款，需大区法务书面批准。"
    negotiation = client.post(
        f"/api/contracts/{contract_id}/negotiations",
        headers=legal,
        json={"status": "客户二次偏离", "customer_feedback": "客户要求维持现有文本。", "internal_note": internal_secret},
    )
    assert negotiation.status_code == 200, negotiation.text

    for round_no in (4, 5):
        next_upload = client.post(
            "/api/contracts/upload",
            headers=legal,
            data={"contract_id": contract_id, "review_round": str(round_no), "contract_type": "租赁", "document_source": "客户"},
            files=[("files", (f"测试租赁合同-R{round_no}.docx", docx_payload(round_two_text), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
        )
        assert next_upload.status_code == 200, next_upload.text
        next_review = client.post(f"/api/contracts/{contract_id}/review", headers=legal)
        assert next_review.status_code == 200, next_review.text
        assert next_review.json()["round_no"] == round_no

    final_detail = client.get(f"/api/contracts/{contract_id}", headers=legal).json()
    assert len(final_detail["versions"]) == 5
    assert final_detail["cumulative_review"]["rounds_completed"] == 5
    assert final_detail["cumulative_review"]["total_opinions"] >= len(final_detail["review_runs"][0]["suggestions"])
    assert client.post(f"/api/contracts/{contract_id}/review", headers=legal).status_code == 400

    latest_detail = client.get(f"/api/contracts/{contract_id}", headers=legal).json()
    latest_run_id = latest_detail["review_runs"][0]["id"]
    blocked_export = client.get(f"/api/contracts/{contract_id}/export?mode=redline", headers=legal)
    assert blocked_export.status_code == 409
    confirmed = client.post(
        f"/api/review-runs/{latest_run_id}/confirm-high-risk",
        headers=legal,
        json={"decision": "确认风险", "comment": "自动化测试：法务已逐项审阅本轮高风险结果。"},
    )
    assert confirmed.status_code == 200, confirmed.text
    assert confirmed.json()["confirmed_count"] > 0
    exported = client.get(f"/api/contracts/{contract_id}/export?mode=redline", headers=legal)
    assert exported.status_code == 200, exported.text
    assert exported.content[:2] == b"PK"
    document = Document(BytesIO(exported.content))
    exported_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "设备租赁合同" in exported_text
    assert any(run.font.highlight_color is not None for paragraph in document.paragraphs for run in paragraph.runs)
    with ZipFile(BytesIO(exported.content)) as package:
        assert "word/comments.xml" in package.namelist()
    clean = client.get(f"/api/contracts/{contract_id}/export?mode=clean", headers=legal)
    clean_document = Document(BytesIO(clean.content))
    clean_text = "\n".join(paragraph.text for paragraph in clean_document.paragraphs)
    assert "设备租赁合同" in clean_text
    assert "修改依据" not in clean_text
    assert internal_secret not in clean_text
    assert clean_text.count(accepted_text) == 1
    internal = client.get(f"/api/contracts/{contract_id}/export?mode=internal", headers=legal)
    assert internal.status_code == 200
    internal_document = Document(BytesIO(internal.content))
    internal_text = "\n".join(paragraph.text for paragraph in internal_document.paragraphs)
    internal_text += "\n" + "\n".join(cell.text for table in internal_document.tables for row in table.rows for cell in row.cells)
    assert internal_secret in internal_text
    assert "Focused Checks 与具体条款依据" in internal_text
    assert "2025年01版.pdf" in internal_text
    assert "第12.2条" in internal_text
    sales = login(client, "sales.owner")
    assert client.get(f"/api/contracts/{contract_id}/export?mode=internal", headers=sales).status_code == 403
    consistency = client.get(f"/api/contracts/{contract_id}/consistency", headers=legal)
    assert consistency.status_code == 200
    assert consistency.json()["versions_checked"] == 5
    redacted = client.get(f"/api/contracts/{contract_id}/redacted-preview", headers=legal)
    assert redacted.status_code == 200


def test_reports_knowledge_and_audit(client: TestClient):
    admin = login(client, "system.admin")
    knowledge = client.get("/api/knowledge", headers=admin).json()
    assert any(item["entry_type"] == "标准条款" for item in knowledge)
    clauses = client.get("/api/standard-clauses?contract_type=租赁", headers=admin)
    assert clauses.status_code == 200
    assert len(clauses.json()) >= 50
    report = client.get("/api/reports/overview", headers=admin)
    assert report.status_code == 200
    assert report.json()["kpis"]["contracts"] >= 4
    audit = client.get("/api/audit-logs", headers=admin)
    assert audit.status_code == 200
    assert any(item["action"] == "完成合同轮次审核并回写原文件" for item in audit.json())
