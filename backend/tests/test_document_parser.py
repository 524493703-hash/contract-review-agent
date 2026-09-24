from __future__ import annotations

from docx import Document

from app.services.document_parser import _is_readable_contract_text, extract_fields
from app.services.review_engine import classify_template


def test_field_extraction_prefers_strong_contract_labels():
    text = """DDBB叉车租赁（框架）合同
合同编号：ZL-2025-001
甲方（承租方）：DDBB物流股份有限公司
合同总金额：人民币 12.5 万元
付款方式：验收合格且收到发票后30日内付款。
"""
    fields = extract_fields(text, "客户版待偏离.docx")
    assert fields["contract_no"] == "ZL-2025-001"
    assert fields["name"] == "DDBB叉车租赁（框架）合同"
    assert fields["customer"] == "DDBB物流股份有限公司"
    assert fields["amount"] == 125000
    assert fields["contract_type"] == "租赁"


def test_readability_accepts_english_purchase_order_and_rejects_control_gibberish():
    english = "Purchase Order Document Number 9200415067 Supplier Linde China " * 12
    gibberish = ("abCD" + "\x00\x07\x11") * 80
    assert _is_readable_contract_text(english)
    assert not _is_readable_contract_text(gibberish)


def test_extract_fields_repairs_ocr_dropped_www_dot():
    fields = extract_fields("供应商质量框架协议，详细信息见网站：wwwVoSS.net，", "PO.pdf")
    assert fields["website_terms_url"] == "https://www.voss.net"


def test_standard_template_and_deviation_classification():
    standard = "林德（中国）叉车有限公司设备销售、交货和保修一般性条款和条件\n1 总则"
    assert classify_template(standard) == "标准合同"
    assert classify_template(standard + "\n[PDF批注 第1页 Highlight] 客户要求修改") == "标准合同（有偏离）"
    assert classify_template("客户自拟叉车租赁框架合同") == "客户版本"


def test_word_comments_are_structured_and_do_not_pollute_contract_body(tmp_path):
    document = Document()
    paragraph = document.add_paragraph("设备租赁合同正文，验收后30日付款。" * 8)
    document.add_comment(paragraph.runs, "客户要求改为60日。", author="客户王经理", initials="客户")
    source = tmp_path / "commented.docx"
    document.save(source)

    from app.services.document_parser import parse_document

    parsed = parse_document(source, source.name)
    assert parsed.status == "已解析"
    assert "客户要求改为60日" not in parsed.text
    assert parsed.annotations[0]["author_name"] == "客户王经理"
    assert parsed.annotations[0]["comment_text"] == "客户要求改为60日。"


def test_docx_parser_preserves_paragraph_table_order_and_source_page(tmp_path, monkeypatch):
    from docx.enum.text import WD_BREAK
    from app.services import document_parser

    document = Document()
    document.add_paragraph("第一条 正文条款")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "未税"
    table.cell(0, 1).text = "含税"
    page = document.add_paragraph()
    page.add_run().add_break(WD_BREAK.PAGE)
    document.add_paragraph("第二条 下一页条款")
    source = tmp_path / "ordered.docx"
    document.save(source)
    monkeypatch.setattr(document_parser, "_rendered_docx_pages", lambda _: [])

    parsed = document_parser.parse_document(source, source.name)
    assert parsed.text.index("第一条") < parsed.text.index("未税 | 含税") < parsed.text.index("第二条")
    second_offset = parsed.text.index("第二条")
    second_map = next(item for item in parsed.source_map if item["start_offset"] <= second_offset < item["end_offset"])
    assert second_map["page_no"] == 2
    assert second_map["source_file"] == source.name
